"""Comparison engine for two sister-vessel PMS Job Status extracts.

compare(df_a, df_b) returns a dict of DataFrames; build_excel(result) returns .xlsx bytes.
"""
import io
import re
import difflib
import html

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# bump whenever compare() output changes, so cached results in the app are recomputed
ENGINE_VERSION = "2.1"

KEY = ["Machinery Location", "Sub Component Location", "Job Code"]
REQUIRED = ["Vessel", "Function", "Machinery Location", "Sub Component Location", "Job Code",
            "Title", "Description", "Frequency", "Performing Rank"]
SHOW_COLS = ["Function", "Title", "Frequency", "Performing Rank", "Maker", "Model", "Job Source"]


def norm(s):
    if pd.isna(s):
        return ""
    s = str(s).replace("&amp;", "&").replace("\xa0", " ")
    return re.sub(r"\s+", " ", s).strip()


def norm_rank(s):
    # order-insensitive set of ranks; handles ',' and '|' separators
    parts = re.split(r"[,|]", norm(s))
    return " | ".join(sorted(p.strip() for p in parts if p.strip()))


def read_extract(src):
    """Read a Job Status extract (path or file-like). Uses the first sheet."""
    df = pd.read_excel(src, dtype={"Job Code": str})
    missing = [c for c in REQUIRED if c not in df.columns]
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(missing)}")
    return df


def prepare(df):
    df = df.copy()
    for c in ["Maker", "Model", "Job Source", "Verifying Rank", "Critical", "E-Form", "Job Status"]:
        if c not in df.columns:
            df[c] = ""
    df["Job Code"] = df["Job Code"].astype(str).str.replace(r"\.0$", "", regex=True)
    # always treat the compared columns as text (an all-blank column, e.g. Critical, would otherwise be float);
    # is_string_dtype also covers pandas 3's dedicated "str" dtype
    text_cols = set(REQUIRED) | {"Maker", "Model", "Job Source", "Verifying Rank", "Critical", "E-Form", "Job Status"}
    for c in df.columns:
        if c in text_cols or pd.api.types.is_object_dtype(df[c]) or pd.api.types.is_string_dtype(df[c]):
            df[c] = df[c].map(norm).astype(object)
    # some extracts repeat the machinery as a prefix in the sub-component path
    # ("Auxiliary Boiler#1 > Auxiliary Boiler - General#1"); strip it so rows still match
    df["Original Sub Component"] = df["Sub Component Location"]
    prefix = df["Machinery Location"] + " > "
    has_prefix = [s.startswith(p) for s, p in zip(df["Sub Component Location"], prefix)]
    df["Prefixed Sub Component"] = has_prefix
    df.loc[has_prefix, "Sub Component Location"] = [
        s[len(p):] for s, p, h in zip(df["Sub Component Location"], prefix, has_prefix) if h]
    # occurrence index so duplicate keys pair up 1:1 instead of cross-joining
    df["Occ"] = df.groupby(KEY).cumcount() + 1
    return df


def word_diff(x, y):
    """Similarity %, words only in x, words only in y."""
    wa, wb = x.split(), y.split()
    sm = difflib.SequenceMatcher(None, wa, wb, autojunk=False)
    removed, added = [], []
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op in ("delete", "replace"):
            removed.append(" ".join(wa[i1:i2]))
        if op in ("insert", "replace"):
            added.append(" ".join(wb[j1:j2]))
    return round(sm.ratio() * 100, 1), " ... ".join(removed)[:1500], " ... ".join(added)[:1500]


def diff_html(x, y):
    """Two HTML strings with removed (red) / added (green) words highlighted."""
    wa, wb = x.split(), y.split()
    sm = difflib.SequenceMatcher(None, wa, wb, autojunk=False)
    left, right = [], []
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        ta, tb = html.escape(" ".join(wa[i1:i2])), html.escape(" ".join(wb[j1:j2]))
        if op == "equal":
            left.append(ta)
            right.append(tb)
        else:
            if ta:
                left.append(f'<span style="background:#ffd6d6;color:#8a0000">{ta}</span>')
            if tb:
                right.append(f'<span style="background:#d4f5d4;color:#005a00">{tb}</span>')
    return " ".join(left), " ".join(right)


def compare(df_a, df_b, name_a=None, name_b=None):
    a, b = prepare(df_a), prepare(df_b)
    NA = name_a or (a["Vessel"].iloc[0] if len(a) else "Vessel A")
    NB = name_b or (b["Vessel"].iloc[0] if len(b) else "Vessel B")
    if NA == NB:
        NA, NB = f"{NA} (A)", f"{NB} (B)"
    sa, sb = f" [{NA}]", f" [{NB}]"

    def count_compare(cols):
        ca = a.groupby(cols).size().rename(NA)
        cb = b.groupby(cols).size().rename(NB)
        r = pd.concat([ca, cb], axis=1).fillna(0).astype(int).reset_index()
        r["Difference (A-B)"] = r[NA] - r[NB]
        r["Status"] = [
            f"Only in {NA}" if y == 0 else f"Only in {NB}" if x == 0 else "Match" if x == y else "Count differs"
            for x, y in zip(r[NA], r[NB])]
        r["_m"] = r["Status"].eq("Match")
        r["_d"] = -r["Difference (A-B)"].abs()
        return r.sort_values(["_m", "_d"]).drop(columns=["_m", "_d"]).reset_index(drop=True)

    func_cmp = count_compare(["Function"])
    mach_cmp = count_compare(["Function", "Machinery Location"])
    sub_cmp = count_compare(["Machinery Location", "Sub Component Location"])

    m = a.merge(b, on=KEY + ["Occ"], how="outer", suffixes=(sa, sb), indicator=True)
    both = m[m["_merge"] == "both"].copy()

    def side(df, s):
        out = df[KEY + [c + s for c in SHOW_COLS]].copy()
        out.columns = KEY + SHOW_COLS
        return out

    def relocated(src, other):
        o = (other.groupby(["Machinery Location", "Job Code"])["Sub Component Location"]
             .apply(lambda x: "; ".join(sorted(set(x)))).rename("Sub Component(s) in other vessel"))
        r = src.merge(o, on=["Machinery Location", "Job Code"], how="left")
        r["Likely cause"] = [
            "Same job on same machinery, different sub-component" if isinstance(x, str) and x
            else "Job missing on this machinery" for x in r["Sub Component(s) in other vessel"]]
        return r

    only_a = relocated(side(m[m["_merge"] == "left_only"], sa), b)
    only_b = relocated(side(m[m["_merge"] == "right_only"], sb), a)

    codes_a, codes_b = set(a["Job Code"]), set(b["Job Code"])
    code_missing = pd.concat([
        a[~a["Job Code"].isin(codes_b)].groupby(["Job Code", "Title"]).size().reset_index(name="Occurrences").assign(Present=f"Only in {NA}"),
        b[~b["Job Code"].isin(codes_a)].groupby(["Job Code", "Title"]).size().reset_index(name="Occurrences").assign(Present=f"Only in {NB}"),
    ], ignore_index=True)

    def field_diff(field, normfn=None):
        fa, fb = both[field + sa], both[field + sb]
        if normfn:
            fa, fb = fa.map(normfn), fb.map(normfn)
        d = both[fa != fb]
        cols = ["Function" + sa] + KEY + ([] if field == "Title" else ["Title" + sa]) + [field + sa, field + sb]
        ren = {"Function" + sa: "Function"}
        if field != "Title":
            ren["Title" + sa] = "Title"
        return d[cols].rename(columns=ren).reset_index(drop=True)

    freq_diff = field_diff("Frequency")
    rank_diff = field_diff("Performing Rank", norm_rank)
    desc_diff = field_diff("Description", lambda s: norm(s).lower())
    if len(desc_diff):
        res = [word_diff(x, y) for x, y in zip(desc_diff["Description" + sa], desc_diff["Description" + sb])]
        desc_diff["Similarity %"] = [r[0] for r in res]
        desc_diff[f"Text only in {NA}"] = [r[1] for r in res]
        desc_diff[f"Text only in {NB}"] = [r[2] for r in res]
        desc_diff = desc_diff.sort_values("Similarity %").reset_index(drop=True)
    else:
        desc_diff["Similarity %"] = []

    other_parts = []
    for f, fn in [("Title", None), ("Verifying Rank", norm_rank), ("Critical", None), ("Job Source", None), ("E-Form", None)]:
        d = field_diff(f, fn)
        if f == "Title":
            d.insert(4, "Title", d[f + sa])
        d = d.rename(columns={f + sa: f"Value{sa}", f + sb: f"Value{sb}"})
        d.insert(0, "Field", f)
        other_parts.append(d)
    other = pd.concat(other_parts, ignore_index=True)

    def mm(df):
        return df.groupby("Machinery Location").apply(
            lambda g: "; ".join(sorted(set(f"{x} / {y}" for x, y in zip(g["Maker"], g["Model"])))), include_groups=False)

    mmc = pd.concat([mm(a).rename(f"Maker/Model{sa}"), mm(b).rename(f"Maker/Model{sb}")], axis=1).fillna("").reset_index()
    mm_diff = mmc[mmc.iloc[:, 1] != mmc.iloc[:, 2]].reset_index(drop=True)

    freq_pattern = (freq_diff.groupby(["Frequency" + sa, "Frequency" + sb]).size()
                    .reset_index(name="Jobs").sort_values("Jobs", ascending=False).reset_index(drop=True))
    rank_pattern = (rank_diff.groupby(["Performing Rank" + sa, "Performing Rank" + sb]).size()
                    .reset_index(name="Jobs").sort_values("Jobs", ascending=False).reset_index(drop=True))

    def per_mach(df, name):
        return df.groupby("Machinery Location").size().rename(name)

    crit = other[other["Field"] == "Critical"]
    hot = pd.concat([
        per_mach(only_a, f"Jobs only in {NA}"), per_mach(only_b, f"Jobs only in {NB}"),
        per_mach(freq_diff, "Frequency diff"), per_mach(rank_diff, "Perf. Rank diff"),
        per_mach(desc_diff, "Description diff"), per_mach(crit, "Critical diff"),
    ], axis=1).fillna(0).astype(int)
    hot["Total discrepancies"] = hot.sum(axis=1)
    fmap = pd.concat([a, b]).drop_duplicates("Machinery Location").set_index("Machinery Location")["Function"]
    hot = hot.sort_values("Total discrepancies", ascending=False).reset_index()
    hot.insert(1, "Function", hot["Machinery Location"].map(fmap))

    # ---- Critical job review ----
    # Criticality should be set per job; a machinery / sub-component where (nearly) every job
    # is critical usually means the flag was applied at equipment level by mistake.
    def crit_stats(df, cols):
        g = df.assign(_c=df["Critical"].str.upper().eq("C")).groupby(cols)["_c"].agg(["size", "sum"])
        return g.rename(columns={"size": "Jobs", "sum": "Critical"})

    def crit_review(cols, min_jobs):
        ga, gb = crit_stats(a, cols), crit_stats(b, cols)
        r = ga.add_suffix(f" - A").join(gb.add_suffix(f" - B"), how="outer").fillna(0).astype(int)
        r["% Critical - A"] = (100 * r["Critical - A"] / r["Jobs - A"].where(r["Jobs - A"] > 0)).round(0).fillna(0)
        r["% Critical - B"] = (100 * r["Critical - B"] / r["Jobs - B"].where(r["Jobs - B"] > 0)).round(0).fillna(0)
        r = r[(r["Critical - A"] > 0) | (r["Critical - B"] > 0)].reset_index()

        def flag(x):
            notes = []
            for t in ("A", "B"):
                if x[f"Jobs - {t}"] >= min_jobs and x[f"% Critical - {t}"] >= 50:
                    notes.append(f"{t}: {x[f'% Critical - {t}']:.0f}% of jobs critical – check if flag applied at equipment level")
            if x["Critical - A"] != x["Critical - B"]:
                notes.append(f"Critical count differs ({x['Critical - A']} vs {x['Critical - B']})")
            return "; ".join(notes) or "OK"
        r["Review note"] = r.apply(flag, axis=1)
        r["_o"] = r["Review note"].str.contains("% of jobs critical")
        r["_s"] = r["Review note"].ne("OK")
        r["_d"] = (r["Critical - A"] - r["Critical - B"]).abs()
        return (r.sort_values(["_o", "_s", "_d"], ascending=False)
                .drop(columns=["_o", "_s", "_d"]).reset_index(drop=True))

    crit_mach = crit_review(["Machinery Location"], 5)
    crit_mach.insert(1, "Function", crit_mach["Machinery Location"].map(
        pd.concat([a, b]).drop_duplicates("Machinery Location").set_index("Machinery Location")["Function"]))
    crit_sub = crit_review(["Machinery Location", "Sub Component Location"], 3)

    ca, cb = crit_stats(a, ["Title"]), crit_stats(b, ["Title"])
    crit_title = ca.add_suffix(" - A").join(cb.add_suffix(" - B"), how="outer").fillna(0).astype(int).reset_index()
    crit_title = crit_title[(crit_title["Critical - A"] > 0) | (crit_title["Critical - B"] > 0)]
    crit_title["Difference (B-A)"] = crit_title["Critical - B"] - crit_title["Critical - A"]
    crit_title = crit_title.sort_values("Difference (B-A)", key=abs, ascending=False).reset_index(drop=True)

    ca_jobs = a[a["Critical"].str.upper().eq("C")][["Function"] + KEY + ["Title", "Frequency"]].assign(Vessel=NA)
    cb_jobs = b[b["Critical"].str.upper().eq("C")][["Function"] + KEY + ["Title", "Frequency"]].assign(Vessel=NB)
    crit_jobs = pd.concat([ca_jobs, cb_jobs], ignore_index=True)
    crit_jobs = crit_jobs[["Vessel"] + [c for c in crit_jobs.columns if c != "Vessel"]]

    naming = pd.concat([
        a[a["Prefixed Sub Component"]].assign(Vessel=NA),
        b[b["Prefixed Sub Component"]].assign(Vessel=NB),
    ])[["Vessel", "Function", "Machinery Location", "Original Sub Component", "Sub Component Location", "Job Code", "Title"]].rename(
        columns={"Original Sub Component": "Sub Component (as in extract)",
                 "Sub Component Location": "Sub Component (normalised)"}).reset_index(drop=True)

    summary = pd.DataFrame([
        ("Total jobs", len(a), len(b)),
        ("Functions", a["Function"].nunique(), b["Function"].nunique()),
        ("Machinery Locations", a["Machinery Location"].nunique(), b["Machinery Location"].nunique()),
        ("Sub Component Locations (machinery+sub)", a.groupby(["Machinery Location", "Sub Component Location"]).ngroups,
         b.groupby(["Machinery Location", "Sub Component Location"]).ngroups),
        ("Distinct Job Codes", len(codes_a), len(codes_b)),
        ("Duplicate job rows (same machinery+sub+code)", int((a["Occ"] > 1).sum()), int((b["Occ"] > 1).sum())),
        ("Sub-component names with duplicated machinery prefix", int(a["Prefixed Sub Component"].sum()), int(b["Prefixed Sub Component"].sum())),
        ("Job Status values", ", ".join(sorted(set(a["Job Status"]) - {""})), ", ".join(sorted(set(b["Job Status"]) - {""}))),
    ], columns=["Metric", NA, NB])
    summary[NA] = summary[NA].astype(str)
    summary[NB] = summary[NB].astype(str)

    findings = pd.DataFrame([
        ("Sub-component naming", f"Names with duplicated machinery prefix ({NB})", int(b["Prefixed Sub Component"].sum()), "Naming Issue"),
        ("Sub-component naming", f"Names with duplicated machinery prefix ({NA})", int(a["Prefixed Sub Component"].sum()), "Naming Issue"),
        ("Counts", "Functions with count difference", int((func_cmp["Status"] != "Match").sum()), "Function Count"),
        ("Counts", "Machinery with count difference", int((mach_cmp["Status"] == "Count differs").sum()), "Machinery Count"),
        ("Counts", f"Machinery only in {NA}", int((mach_cmp["Status"] == f"Only in {NA}").sum()), "Machinery Count"),
        ("Counts", f"Machinery only in {NB}", int((mach_cmp["Status"] == f"Only in {NB}").sum()), "Machinery Count"),
        ("Counts", "Sub-components with count difference", int((sub_cmp["Status"] == "Count differs").sum()), "SubComponent Count"),
        ("Counts", f"Sub-components only in {NA}", int((sub_cmp["Status"] == f"Only in {NA}").sum()), "SubComponent Count"),
        ("Counts", f"Sub-components only in {NB}", int((sub_cmp["Status"] == f"Only in {NB}").sum()), "SubComponent Count"),
        ("Jobs", "Jobs matched (Machinery + Sub-component + Job Code)", len(both), ""),
        ("Jobs", f"Jobs only in {NA}", len(only_a), "Jobs Only A"),
        ("Jobs", f"   of which same job on same machinery, other sub-component", int(only_a["Likely cause"].str.startswith("Same").sum()), "Jobs Only A"),
        ("Jobs", f"Jobs only in {NB}", len(only_b), "Jobs Only B"),
        ("Jobs", f"   of which same job on same machinery, other sub-component", int(only_b["Likely cause"].str.startswith("Same").sum()), "Jobs Only B"),
        ("Jobs", "Job Codes not used at all on the other vessel", len(code_missing), "JobCode Missing"),
        ("Job details", "Frequency differs", len(freq_diff), "Frequency Diff"),
        ("Job details", "Performing Rank differs (order ignored)", len(rank_diff), "PerfRank Diff"),
        ("Job details", "Description differs", len(desc_diff), "Description Diff"),
        ("Critical", f"Critical jobs on {NA}", int(a["Critical"].str.upper().eq("C").sum()), "Critical Jobs"),
        ("Critical", f"Critical jobs on {NB}", int(b["Critical"].str.upper().eq("C").sum()), "Critical Jobs"),
        ("Critical", "Machinery with ≥50% of jobs critical (possible over-flagging)",
         int(crit_mach["Review note"].str.contains("% of jobs critical").sum()), "Critical by Machinery"),
        ("Job details", "Title differs", int((other["Field"] == "Title").sum()), "Other Field Diff"),
        ("Job details", "Verifying Rank differs", int((other["Field"] == "Verifying Rank").sum()), "Other Field Diff"),
        ("Job details", "Critical flag differs", int((other["Field"] == "Critical").sum()), "Other Field Diff"),
        ("Job details", "Job Source differs", int((other["Field"] == "Job Source").sum()), "Other Field Diff"),
        ("Job details", "E-Form differs", int((other["Field"] == "E-Form").sum()), "Other Field Diff"),
        ("Equipment", "Machinery with different Maker/Model", len(mm_diff), "Maker Model Diff"),
    ], columns=["Area", "Check", "Count", "Sheet"])

    return {
        "version": ENGINE_VERSION,
        "names": (NA, NB),
        "critical_counts": (int(a["Critical"].str.upper().eq("C").sum()), int(b["Critical"].str.upper().eq("C").sum())),
        "overflagged": int(crit_mach["Review note"].str.contains("% of jobs critical").sum()),
        "summary": summary,
        "findings": findings,
        "Naming Issue": naming,
        "Function Count": func_cmp,
        "Machinery Count": mach_cmp,
        "SubComponent Count": sub_cmp,
        "Jobs Only A": only_a,
        "Jobs Only B": only_b,
        "JobCode Missing": code_missing,
        "Frequency Diff": freq_diff,
        "Freq Patterns": freq_pattern,
        "PerfRank Diff": rank_diff,
        "Rank Patterns": rank_pattern,
        "Description Diff": desc_diff,
        "Other Field Diff": other,
        "Maker Model Diff": mm_diff,
        "Hotspots": hot,
        "Critical by Machinery": crit_mach,
        "Critical by SubComp": crit_sub,
        "Critical by Title": crit_title,
        "Critical Jobs": crit_jobs,
    }


SHEET_ORDER = ["Naming Issue", "Function Count", "Machinery Count", "SubComponent Count", "Jobs Only A", "Jobs Only B",
               "JobCode Missing", "Frequency Diff", "Freq Patterns", "PerfRank Diff", "Rank Patterns",
               "Description Diff", "Other Field Diff", "Critical by Machinery", "Critical by SubComp",
               "Critical by Title", "Critical Jobs", "Maker Model Diff", "Hotspots"]


def build_excel(result, sheets=None, only_mismatch_counts=True):
    """Build the Excel report. `sheets` limits which detail sheets are included."""
    NA, NB = result["names"]
    sheets = sheets or SHEET_ORDER
    hdr = PatternFill("solid", fgColor="1F4E78")
    buf = io.BytesIO()
    summary, findings = result["summary"], result["findings"]
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="Summary", index=False, startrow=2)
        f_row = len(summary) + 6
        findings.to_excel(xw, sheet_name="Summary", index=False, startrow=f_row)
        ws = xw.sheets["Summary"]
        ws["A1"] = f"Job comparison: A = {NA}  vs  B = {NB}"
        ws["A1"].font = Font(bold=True, size=14)
        ws.cell(f_row, 1, "Discrepancy findings").font = Font(bold=True, size=12)
        for row in (3, f_row + 1):
            for c in ws[row]:
                if c.value is not None:
                    c.font, c.fill = Font(bold=True, color="FFFFFF"), hdr
        for name in sheets:
            df = result[name]
            if only_mismatch_counts and name in ("Machinery Count", "SubComponent Count", "Function Count"):
                df = df[df["Status"] != "Match"]
            df = df.copy()
            for c in df.columns:
                if pd.api.types.is_object_dtype(df[c]) or pd.api.types.is_string_dtype(df[c]):
                    df[c] = df[c].astype(str).str[:32000]
            df.to_excel(xw, sheet_name=name[:31], index=False)
            ws = xw.sheets[name[:31]]
            for c in ws[1]:
                c.font, c.fill = Font(bold=True, color="FFFFFF"), hdr
                c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
        for ws in xw.book.worksheets:
            for i, col in enumerate(ws.columns, 1):
                width = max((len(str(c.value)) for c in list(col)[:200] if c.value is not None), default=8)
                ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)
    return buf.getvalue()
