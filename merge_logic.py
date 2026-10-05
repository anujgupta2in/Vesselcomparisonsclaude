"""
Merge Job Description, Maker and Model from a PMS "Job List" export
into a PMS "Job Status" export.

A Job Code can be linked to several machineries (e.g. Job Code 1 is on both
S-Band and X-Band radar), so rows are matched on Job Code + Machinery +
Component, with progressively looser fallbacks:

  1  Code + Machinery + Component   (exact / prefix component match)
  2  Code + Machinery               (only one Job List entry, or all entries agree)
  3  Code + Machinery (ambiguous)   (entries differ -> values joined with " | ", flagged)
  4  Code only                      (code exists on a single machinery only)
  -  Not Found

Can be used from the Streamlit app (app.py) or from the command line:
    python merge_logic.py "<job status file>" "<job list file>" [output.xlsx]
"""

import io
import re
import sys
from pathlib import Path

import pandas as pd

NEW_COLS = ["Description", "Verifying Rank", "Maker", "Model", "Match Level", "Match Note"]

# Final column layout of the output file. Match Level / Match Note are appended at the end.
OUTPUT_ORDER = [
    "Vessel", "Function", "Machinery Location", "Sub Component Location", "Job Code", "Title",
    "Description", "Frequency", "Performing Rank", "Verifying Rank", "Maker", "Model", "Critical",
    "Calculated Due Date", "Due Date", "Next Due", "Completion Date", "Job Status", "Job Action",
    "Remaining Running Hours", "Machinery Running Hours", "Last Done Running Hours",
    "Last Done Date", "CMS Code", "Job Source", "E-Form", "Attachment Indicator",
]
DATE_COLS = ["Calculated Due Date", "Due Date", "Next Due", "Completion Date", "Last Done Date"]
NUMBER_COLS = ["Job Code", "Remaining Running Hours", "Machinery Running Hours", "Last Done Running Hours"]

LEVEL_1 = "1 - Code + Machinery + Component"
LEVEL_2 = "2 - Code + Machinery"
LEVEL_3 = "3 - Code + Machinery (ambiguous - review)"
LEVEL_4 = "4 - Code only"
NOT_FOUND = "Not Found"

# Candidate header names for each logical field (first match wins, case-insensitive)
JOB_LIST_FIELDS = {
    "code": ["Job Code"],
    "machinery": ["Machinery", "Machinery Location"],
    "component": ["Component", "Sub Component Location", "Sub Component"],
    "description": ["Description", "Job Description"],
    "maker": ["Maker"],
    "model": ["Model"],
    "verifying": ["Verifying Rank"],
    "title": ["Job Title", "Title"],
}
JOB_STATUS_FIELDS = {
    "code": ["Job Code"],
    "machinery": ["Machinery Location", "Machinery"],
    "component": ["Sub Component Location", "Component", "Sub Component"],
    "title": ["Title", "Job Title"],
}
REQUIRED_LIST = ["code", "machinery", "description", "maker", "model"]
REQUIRED_STATUS = ["code", "machinery"]

EXCEL_CELL_LIMIT = 32767


# --------------------------------------------------------------------------- #
# Reading
# --------------------------------------------------------------------------- #
def read_table(source, name: str = "") -> pd.DataFrame:
    """Read a CSV or Excel file (path or uploaded file object) as all-text."""
    name = (name or str(getattr(source, "name", source))).lower()
    data = source.read() if hasattr(source, "read") else Path(source).read_bytes()

    if name.endswith((".xlsx", ".xlsm", ".xls")):
        df = pd.read_excel(io.BytesIO(data), dtype=str)
    else:
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                df = pd.read_csv(io.BytesIO(data), dtype=str, encoding=enc)
                break
            except UnicodeDecodeError:
                continue
    return df.fillna("")


def find_columns(df: pd.DataFrame, fields: dict, required: list, label: str) -> dict:
    lookup = {str(c).strip().lower(): c for c in df.columns}
    found = {}
    for key, options in fields.items():
        for opt in options:
            if opt.lower() in lookup:
                found[key] = lookup[opt.lower()]
                break
    missing = [fields[k][0] for k in required if k not in found]
    if missing:
        raise ValueError(f"{label} file is missing column(s): {', '.join(missing)}")
    return found


# --------------------------------------------------------------------------- #
# Normalisation
# --------------------------------------------------------------------------- #
# Job Status spells out locations (Starboard1, Fwd-Port1) where Job List
# abbreviates them (S1, F-P1). Applied identically to both files.
_LOCATION_WORDS = [
    (r"Starboard|Stbd", "S"),
    (r"Port", "P"),
    (r"Forward|Fwd", "F"),
    (r"Aft", "A"),
]


def norm_code(v: str) -> str:
    v = str(v).strip()
    return v[:-2] if v.endswith(".0") else v


def norm_machinery(v: str) -> str:
    v = str(v).strip()
    for pattern, abbr in _LOCATION_WORDS:
        # only when the word sits at a location-suffix position: followed by digit, '-', space or end
        v = re.sub(rf"({pattern})(?=[\d\-\s]|$)", abbr, v)
    return re.sub(r"\s+", "", v).lower()


def norm_component(v: str) -> str:
    v = str(v).split(">")[-1]  # "Parent#1 > Child#1" -> "Child#1"
    v = re.sub(r"#\d+\s*$", "", v.strip())
    return re.sub(r"\s+", " ", v).strip().lower()


# --------------------------------------------------------------------------- #
# Matching
# --------------------------------------------------------------------------- #
def _component_matches(cands: list, comp: str) -> list:
    """Best component matches: exact first, else the longest Job List prefix."""
    if not comp:
        return []
    exact = [c for c in cands if c["_comp"] == comp]
    if exact:
        return exact
    prefix = [c for c in cands if c["_comp"] and comp.startswith(c["_comp"])]
    if prefix:
        best = max(len(c["_comp"]) for c in prefix)
        return [c for c in prefix if len(c["_comp"]) == best]
    return []


def _combine(cands: list, field: str, sep: str) -> str:
    seen = []
    for c in cands:
        v = str(c[field]).strip()
        if v and v not in seen:
            seen.append(v)
    return sep.join(seen)


def _values(cands: list) -> tuple:
    return (
        _combine(cands, "description", "\n\n---------- next entry ----------\n\n"),
        _combine(cands, "verifying", " | "),
        _combine(cands, "maker", " | "),
        _combine(cands, "model", " | "),
    )


def _all_same(cands: list) -> bool:
    return len({(c["description"].strip(), c["maker"].strip(), c["model"].strip()) for c in cands}) == 1


def _find_critical_column(df: pd.DataFrame):
    """The Job Status export has no header for the Critical flag: a blank column holding 'C'."""
    for c in df.columns:
        if str(c).strip().lower() == "critical":
            return c
    for c in df.columns:
        if str(c).startswith("Unnamed:"):
            vals = set(df[c].str.strip().str.upper()) - {""}
            if vals and vals <= {"C", "Y", "YES"}:
                return c
    return None


def arrange_columns(df: pd.DataFrame, s_cols: dict) -> pd.DataFrame:
    """Put columns in OUTPUT_ORDER (case-insensitive header match); missing ones left blank."""
    df = df.copy()
    crit = _find_critical_column(df)
    if crit is not None and crit != "Critical":
        df = df.rename(columns={crit: "Critical"})
    lookup = {str(c).strip().lower(): c for c in df.columns}
    # Job Status may use alternative headers for the key fields
    aliases = {"machinery location": s_cols.get("machinery"),
               "sub component location": s_cols.get("component"),
               "title": s_cols.get("title")}

    out = pd.DataFrame(index=df.index)
    for col in OUTPUT_ORDER:
        src = lookup.get(col.lower()) or aliases.get(col.lower())
        out[col] = df[src] if src is not None else ""
    out["Match Level"] = df["Match Level"]
    out["Match Note"] = df["Match Note"]
    return out


def merge(job_status: pd.DataFrame, job_list: pd.DataFrame):
    """Return (merged DataFrame, summary dict, warnings list)."""
    warnings = []
    s_cols = find_columns(job_status, JOB_STATUS_FIELDS, REQUIRED_STATUS, "Job Status")
    l_cols = find_columns(job_list, JOB_LIST_FIELDS, REQUIRED_LIST, "Job List")
    if "component" not in s_cols or "component" not in l_cols:
        warnings.append("Component column not found in one file - matching on Code + Machinery only.")

    # Vessel sanity check
    if "Vessel" in job_status.columns and "Vessel" in job_list.columns:
        vs = set(job_status["Vessel"].str.strip()) - {""}
        vl = set(job_list["Vessel"].str.strip()) - {""}
        if vs and vl and not vs & vl:
            warnings.append(f"Vessel names differ between files: Job Status {sorted(vs)} vs Job List {sorted(vl)}")

    # Index Job List by code and by (code, machinery)
    by_code, by_code_mach = {}, {}
    for _, r in job_list.iterrows():
        rec = {
            "code": norm_code(r[l_cols["code"]]),
            "_mach": norm_machinery(r[l_cols["machinery"]]),
            "_comp": norm_component(r[l_cols["component"]]) if "component" in l_cols else "",
            "machinery": r[l_cols["machinery"]],
            "component": r[l_cols["component"]] if "component" in l_cols else "",
            "description": r[l_cols["description"]],
            "maker": r[l_cols["maker"]],
            "model": r[l_cols["model"]],
            "verifying": r[l_cols["verifying"]] if "verifying" in l_cols else "",
        }
        by_code.setdefault(rec["code"], []).append(rec)
        by_code_mach.setdefault((rec["code"], rec["_mach"]), []).append(rec)

    results = []
    for _, r in job_status.iterrows():
        code = norm_code(r[s_cols["code"]])
        mach = norm_machinery(r[s_cols["machinery"]])
        comp = norm_component(r[s_cols["component"]]) if "component" in s_cols else ""
        cands = by_code_mach.get((code, mach), [])

        if cands:
            comp_hits = _component_matches(cands, comp)
            if comp_hits and _all_same(comp_hits):
                level, use, note = LEVEL_1, comp_hits[:1], ""
            elif len(cands) == 1 or _all_same(cands):
                level, use, note = LEVEL_2, cands[:1], ""
            else:
                use = comp_hits or cands
                level = LEVEL_3
                note = (f"{len(use)} Job List entries for this code/machinery with different values: "
                        + "; ".join(sorted({c['component'] or '(no component)' for c in use})))
        else:
            code_cands = by_code.get(code, [])
            machs = {c["_mach"] for c in code_cands}
            if code_cands and len(machs) == 1:
                level, use = LEVEL_4, code_cands
                note = f"Machinery name differs - Job List has '{code_cands[0]['machinery']}'"
                if not _all_same(use):
                    level = LEVEL_3
            elif code_cands:
                level, use = NOT_FOUND, []
                note = ("Job Code exists on other machinery: "
                        + ", ".join(sorted({c['machinery'] for c in code_cands})))
            else:
                level, use, note = NOT_FOUND, [], "Job Code not in Job List"

        values = _values(use) if use else ("", "", "", "")
        results.append((*values, level, note))

    out = job_status.copy()
    new = pd.DataFrame(results, columns=NEW_COLS, index=out.index)
    for c in NEW_COLS:
        if c in out.columns:  # avoid clashing with an existing column
            out = out.drop(columns=c)
    out = pd.concat([out, new], axis=1)
    out = arrange_columns(out, s_cols)

    summary = new["Match Level"].value_counts().reindex(
        [LEVEL_1, LEVEL_2, LEVEL_3, LEVEL_4, NOT_FOUND], fill_value=0).to_dict()
    return out, summary, warnings


# --------------------------------------------------------------------------- #
# Export
# --------------------------------------------------------------------------- #
def to_excel_bytes(df: pd.DataFrame) -> bytes:
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    from datetime import datetime

    from openpyxl import Workbook
    from openpyxl.styles import Border, Side

    def cell_value(col, v):
        v = ILLEGAL_CHARACTERS_RE.sub("", str(v)).strip()
        if not v:
            return None
        if col in DATE_COLS:
            for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%d-%b-%Y"):
                try:
                    return datetime.strptime(v, fmt)
                except ValueError:
                    pass
        if col in NUMBER_COLS:
            try:
                f = float(v.replace(",", ""))
                return int(f) if f.is_integer() else f
            except ValueError:
                pass
        return v[:EXCEL_CELL_LIMIT]

    wb = Workbook()
    ws = wb.active
    ws.title = "Job Status"
    cols = list(df.columns)

    thin = Side(style="thin", color="8EA9DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="B4C6E7")
    header_font = Font(name="Calibri", size=10, bold=True, color="1F3864")
    body_font = Font(name="Calibri", size=10)
    review_fill = PatternFill("solid", fgColor="FFF2CC")
    missing_fill = PatternFill("solid", fgColor="F8CBAD")
    top_left = Alignment(vertical="top", wrap_text=False)
    wrap_top = Alignment(vertical="top", wrap_text=True)
    center_top = Alignment(horizontal="center", vertical="top")

    ws.append(cols)
    for cell in ws[1]:
        cell.fill, cell.font, cell.border = header_fill, header_font, border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30

    for record in df.itertuples(index=False):
        ws.append([cell_value(c, v) for c, v in zip(cols, record)])

    wide = {"Description": 60, "Title": 40, "Machinery Location": 32, "Sub Component Location": 36,
            "Performing Rank": 28, "Match Level": 34, "Match Note": 50, "Function": 26}
    for i, col in enumerate(cols, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = wide.get(col, max(len(col) * 0.9 + 2, 11))
        if col in DATE_COLS:
            fmt, align = "DD-MM-YYYY", center_top
        elif col in NUMBER_COLS or col == "Critical":
            fmt, align = "General", center_top
        else:
            fmt, align = "@", wrap_top if col == "Description" else top_left
        for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
            cell.font, cell.border, cell.alignment, cell.number_format = body_font, border, align, fmt
            if col == "Match Level" and cell.value == NOT_FOUND:
                cell.fill = missing_fill
            elif col == "Match Level" and cell.value in (LEVEL_3, LEVEL_4):
                cell.fill = review_fill

    ws.freeze_panes = "G2"  # keep Vessel..Title visible while scrolling
    ws.auto_filter.ref = ws.dimensions

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_csv_bytes(df: pd.DataFrame) -> bytes:
    return df.to_csv(index=False).encode("utf-8-sig")


# --------------------------------------------------------------------------- #
# Command line
# --------------------------------------------------------------------------- #
if __name__ == "__main__":
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)
    status_path, list_path = sys.argv[1], sys.argv[2]
    out_path = sys.argv[3] if len(sys.argv) > 3 else str(
        Path(status_path).with_name(Path(status_path).stem + " - with Job Details.xlsx"))

    merged, summary, warns = merge(read_table(status_path), read_table(list_path))
    for w in warns:
        print("WARNING:", w)
    for k, v in summary.items():
        print(f"{k:45s} {v}")
    Path(out_path).write_bytes(
        to_csv_bytes(merged) if out_path.lower().endswith(".csv") else to_excel_bytes(merged))
    print("Saved:", out_path)
